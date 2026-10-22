import os, datetime, tempfile
HERE=os.path.dirname(os.path.realpath(__file__))
__file__=os.path.join(HERE,'test_ranch.py')
exec(open(__file__).read().split('with sync_playwright() as pw:')[0].replace('PORT=8765','PORT=8783'))
# Synthetic Cow-Calf SPA worksheet (same labels and layout as the Texas A&M IRM-SPA Data Entry and SPASum sheets; made-up ranch)
import openpyxl
def xl(d): return (d - datetime.date(1899,12,30)).days
wb=openpyxl.Workbook(); D=wb.active; D.title='Data Entry'
rows=[
 (4,['Farm or Ranch/Unit :','Test Ranch']),(11,['2','Fiscal Year:','Fiscal Year :',2025]),(16,['County :','Wheeler']),
 (25,['8','Size of herd - Breeding Cows :',220,'Head']),
 (31,['    (c) Calving',xl(datetime.date(2025,3,1)),xl(datetime.date(2025,4,30)),60]),
 (32,['    (d) Weaning',xl(datetime.date(2025,9,20)),xl(datetime.date(2025,10,10)),20]),
 (97,['(g) Totals',4000,4000]),(110,['Total Land Leased',4000,4000]),(111,['Total Land Used',4000,4000]),
 (173,['Roughage',20000,100]),(176,['Protein Supplement',50000,250]),(180,['16','Total Raised & Purchased Feed',70000,350]),
 (209,['17c.','Weaned Calves -- Dominant Breed:','Angus','Angus']),
 (246,['(e) Total',150,75000,240000,500,1600,320]),
 (266,['21','Total females exposed at the']),(267,['beginning of the breeding season.',200]),
 (309,['29','Adjusted exposed females including sales, transfers,']),(310,['purchase of pairs and exposed and pregnant females']),(311,['(21-22-23+24-25b+26-27+28).',200]),
 (315,['30a.','Total death loss of exposed females.',4]),
 (319,['31','Replacement heifers exposed for first']),(320,['calves plus purchased replacement heifers']),(321,['and cows to expose or exposed.',30]),
 (327,['33','Number of exposed females that were pregnancy tested.',200]),(329,['34a.','Number of females diagnosed as pregnant.',184]),
 (352,['38a.','Total calves born --']),(353,['include live and dead calving (head).',176]),(368,['40','Total live calves born (head).',172]),
 (384,['43','Total weaned calves.',168]),(386,['44','Average Age at weaning (months).',6.5]),
]
for r,vals in rows:
    for i,v in enumerate(vals): D.cell(r,2+i,v)
S=wb.create_sheet('SPASum')
for i,(k,v) in enumerate([('Average Weaning Weight',500),('Raised/Purchased Feed Cost per cow',150),('Grazing Cost per cow',0),('Total Cost Before Non Calf Revenue Adjusted per cow',900),('Total Cost Noncalf Revenue Adjusted per cwt',210),('Net Income After Withdrawal per cow',400)]):
    S.cell(10+i,1,k); S.cell(10+i,2,v)
I=wb.create_sheet('ISData')
for r,vals in [(20,['Feed Purchase',30000,100,30000]),(24,['   - Protein',25000,100,25000]),(25,['   - Mineral & Additives',5000,100,5000]),(90,['Total Operating Expenses',180000,180000])]:
    for i,v in enumerate(vals): I.cell(r,1+i,v)
tmp=os.path.join(tempfile.mkdtemp(),'Test-SPA.xlsx'); wb.save(tmp)

with sync_playwright() as pw:
    b=pw.chromium.launch(); errs=[]
    c=b.new_context(viewport={'width':1360,'height':1300}); c.route('**/*', route); pg=c.new_page(); pg.on('pageerror', lambda e: errs.append(str(e))); pg.on('dialog', lambda d: d.accept())
    pg.goto(f'http://127.0.0.1:{PORT}/my-ranch/index.html'); pg.wait_for_timeout(700)
    pg.click('#summaryBody [data-example="rolling-plains-154-cows"]'); pg.wait_for_timeout(2000)
    # other costs: defaults from the BCU sheet, net return and break-even shown
    oc=pg.evaluate("MyRanch.herd().oc")
    check('other costs start at the BCU sheet values', oc['land']==165 and oc['labor']==120 and oc['vet']==35 and oc['vehicle']==48, oc)
    pg.evaluate("MyRanch.go('plan')"); pg.wait_for_timeout(800)
    tl=pg.evaluate("document.getElementById('fc_tiles').textContent")
    check('year plan shows net return after other costs and the break-even calf price', 'Net return' in tl and 'of other costs' in tl and 'break-even calf price $' in tl, tl)
    pg.evaluate("MyRanch.go('herd')"); pg.wait_for_timeout(400)
    n0=pg.evaluate("document.getElementById('st_net').textContent")
    pg.evaluate("var e=document.getElementById('oc_labor'); e.value='220'; e.dispatchEvent(new Event('input'))"); pg.wait_for_timeout(500)
    n1=pg.evaluate("document.getElementById('st_net').textContent")
    check('top bar shows net return and it drops when a cost goes up', 'Net return' in pg.evaluate("document.getElementById('st_net').parentElement.textContent") and n0 != n1, (n0, n1))
    pg.evaluate("document.getElementById('oc_reset').click()"); pg.wait_for_timeout(300)
    check('back-to-BCU button says what it did', 'Back to the BCU sheet' in pg.evaluate("document.getElementById('oc_note').textContent") and pg.evaluate("MyRanch.herd().oc.labor")==120)
    pg.evaluate("document.getElementById('oc_reset').click()"); pg.wait_for_timeout(300)
    check('pressing it again says nothing changed', 'already are' in pg.evaluate("document.getElementById('oc_note').textContent"))
    pg.evaluate("document.getElementById('oc_fromBcu').click()"); pg.wait_for_timeout(300)
    check('home-page button explains when there is nothing saved or nothing different', 'home page' in pg.evaluate("document.getElementById('oc_note').textContent"))
    pg.evaluate("var e=document.getElementById('oc_labor'); e.value='220'; e.dispatchEvent(new Event('input'))"); pg.wait_for_timeout(500)
    check('changing a cost moves the net', pg.evaluate("MyRanch.herd().oc.labor")==220 and 'per cow' in pg.evaluate("document.getElementById('oc_note').textContent"))
    # SPA worksheet
    pg.set_input_files('#spaInput', tmp); pg.wait_for_timeout(2500)
    h=pg.evaluate("MyRanch.herd()")
    check('SPA fills cows (exposed), calving start and length, weaning date', h['cows']==200 and h['start1']=='03-01' and h['len1']==60 and h['wean1']=='09-30', (h['cows'],h['start1'],h['len1'],h['wean1']))
    check('SPA fills calving rate, calf loss, cow loss, replacements, records weight', h['calvingRate']==88 and abs(h['calfLoss']-4.5)<0.01 and h['cowLoss']==2 and h['replRate']==15 and h['weanLb']==500, (h['calvingRate'],h['calfLoss'],h['cowLoss'],h['replRate'],h['weanLb']))
    check('SPA non-feed costs go to Other (per exposed female)', h['oc']['other']==round((900-150)*220/200) and h['oc']['labor']==0, h['oc'])
    check('SPA note says what was filled', 'Test-SPA.xlsx' in pg.inner_text('#spaNote') and '2025' in pg.inner_text('#spaNote'))
    pg.evaluate("MyRanch.go('plan')"); pg.wait_for_timeout(1200)
    t=pg.inner_text('#fc_spa')
    check('year plan compares SPA records with My Ranch', pg.is_visible('#fc_spa') and 'SPA records (2025)' in t and 'Calves weaned' in t and '168' in t and 'Protein supplement' in t and '250' in t, t[:300])
    check('worksheet checks flag breeding cows vs exposed, unweighed calves, $0 grazing on leased land', 'differ by more than 10%' in t and 'Only 150 of 168' in t and 'Grazing cost is $0' in t, t)
    check('records show the protein price paid ($25,000 for 25 t = $1,000/t) and offer to use it', 'Protein supplement price' in t and '$1,000' in t, t)
    pg.evaluate("document.getElementById('spaUsePrice').click()"); pg.wait_for_timeout(1200)
    check('use the price paid sets the supplement price', pg.evaluate("MyRanch.herd().supPrice")==1000)
    st=pg.evaluate("MyRanch.state().spa")
    check('SPA records saved with the ranch', st and st['weaned']==168 and st['protLb']==50000 and st['roughLb']==20000)
    pg.click('#spaClear'); pg.wait_for_timeout(500)
    check('SPA records can be removed', not pg.is_visible('#fc_spa') and pg.evaluate("MyRanch.state().spa")==None)
    pg.evaluate("MyRanch.go('herd')"); pg.wait_for_timeout(300)
    bad=os.path.join(tempfile.mkdtemp(),'notes.xlsx'); w2=openpyxl.Workbook(); w2.active.title='Sheet1'; w2.save(bad)
    pg.set_input_files('#spaInput', bad); pg.wait_for_timeout(1500)
    check('a workbook that is not the SPA worksheet is refused with a message', 'No “Data Entry” sheet' in pg.inner_text('#spaNote'))
    check('no JS errors', not errs, errs)
print('\n%d failures' % len(fails)); sys.exit(1 if fails else 0)
